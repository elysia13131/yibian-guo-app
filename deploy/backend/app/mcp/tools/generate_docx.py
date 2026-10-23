import os
import re
import logging
import tempfile
from pathlib import Path
from typing import Dict, Any, Optional

import httpx

logger = logging.getLogger(__name__)

_IMAGE_PATTERN = re.compile(r'!\[([^\]]*)\]\(([^)]+)\)')

_OUTPUT_BASE = Path(__file__).resolve().parent.parent.parent / "_output"
_API_IMAGE_PREFIX = "/api/v1/image/output/"


async def _resolve_image_src(src: str) -> Optional[str]:
    """将图片源（远程 URL、API 路径或本地路径）解析为临时文件路径"""
    if src.startswith(('http://', 'https://')):
        return await _download_image(src)
    if src.startswith('file:///'):
        src = src[8:]
    if src.startswith(_API_IMAGE_PREFIX):
        local_path = (_OUTPUT_BASE / src[len(_API_IMAGE_PREFIX):]).resolve()
        if local_path.exists():
            return str(local_path)
    p = Path(src)
    if p.exists():
        return str(p)
    return None


def _strip_markdown(text: str) -> str:
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'^[\*\-]\s+', '', text, flags=re.MULTILINE)
    text = text.replace('*', '')
    text = re.sub(r'`(.+?)`', r'\1', text)
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
    text = text.replace('\u2003', '  ').replace('\u00a0', ' ')
    text = _IMAGE_PATTERN.sub(r'[图片: \1]', text)
    return text


async def _download_image(url: str) -> Optional[str]:
    """下载远程图片到临时文件，返回路径"""
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            resp = await client.get(url)
            resp.raise_for_status()
            suffix = Path(url.split('?')[0]).suffix or '.png'
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
            tmp.write(resp.content)
            tmp.close()
            return tmp.name
    except Exception as e:
        logger.warning(f"图片下载失败 {url}: {e}")
        return None


async def generate_docx(params: Dict[str, Any]) -> str:
    """将 Markdown 文本内容生成为 Word (DOCX) 文件，支持嵌入远程图片。

    Args:
        params: {
            "title": "文档标题",
            "content": "Markdown 格式的正文内容（支持 ![alt](url) 图片）",
            "filename": "可选，输出文件名（不含扩展名）"
        }
    Returns:
        生成的 DOCX 文件路径
    """
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn

    title = params.get("title", "文档")
    content = params.get("content", "")
    filename = params.get("filename", re.sub(r'[\/\\:*?"<>|]', '_', title))

    output_dir = Path(__file__).resolve().parent.parent.parent / "_output"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{filename}.docx"

    doc = Document()
    style = doc.styles['Normal']
    style.font.size = Pt(11)
    style.font.name = '宋体'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    p = doc.add_paragraph()
    p.alignment = 1
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(18)

    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 检测图片行
        img_match = _IMAGE_PATTERN.search(stripped)
        if img_match:
            alt_text = img_match.group(1)
            img_src = img_match.group(2)
            img_path = await _resolve_image_src(img_src)
            if img_path:
                is_temp = img_src.startswith(('http://', 'https://'))
                try:
                    doc.add_picture(img_path, width=Inches(5))
                    if alt_text:
                        cap = doc.add_paragraph()
                        cap.alignment = 1
                        cap_run = cap.add_run(alt_text)
                        cap_run.font.size = Pt(9)
                        cap_run.font.italic = True
                except Exception as e:
                    logger.warning(f"DOCX 嵌入图片失败: {e}")
                    doc.add_paragraph(f"[图片: {alt_text}]")
                finally:
                    if is_temp:
                        os.unlink(img_path)
            else:
                doc.add_paragraph(f"[图片: {alt_text}]")
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith('#'):
            level = len(stripped.split(' ')[0])
            text = stripped.lstrip('#').strip()
            doc.add_heading(text, level=min(level, 4))

        elif stripped.startswith('|') and stripped.endswith('|'):
            sep_line = lines[i + 1].strip() if i + 1 < len(lines) else ''
            if re.match(r'^\|[-\s:|]+\|$', sep_line):
                header = [c.strip() for c in stripped.strip('|').split('|')]
                rows = []
                j = i + 2
                while j < len(lines):
                    rl = lines[j].strip()
                    if rl.startswith('|') and rl.endswith('|'):
                        cells = [c.strip() for c in rl.strip('|').split('|')]
                        rows.append(cells)
                        j += 1
                    else:
                        break
                table = doc.add_table(rows=1 + len(rows), cols=len(header))
                table.style = 'Table Grid'
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ci, h in enumerate(header):
                    cell = table.rows[0].cells[ci]
                    cell.text = h
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                for ri, row in enumerate(rows):
                    for ci, cell_text in enumerate(row):
                        table.rows[ri + 1].cells[ci].text = cell_text
                doc.add_paragraph()
                i = j
                continue

        elif stripped.startswith('- ') or stripped.startswith('* '):
            text = re.sub(r'^[\*\-]\s+', '', stripped)
            doc.add_paragraph(_strip_markdown(text), style='List Bullet')

        else:
            doc.add_paragraph(_strip_markdown(stripped))

        i += 1

    doc.save(str(output_path))
    logger.info(f"DOCX 已生成: {output_path}")
    return str(output_path)
