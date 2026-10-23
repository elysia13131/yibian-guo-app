import logging
import os
import io
import tempfile
import json
from typing import List, Optional, Dict, Any
from PIL import Image, ImageDraw, ImageFont

from app.services.somark_request_manager import somark_request_manager

logger = logging.getLogger(__name__)


class ExpressionService:

    COLLAGE_MAX_WIDTH = 2400
    COLLAGE_PADDING = 40
    LABEL_HEIGHT = 60
    FONT_SIZE = 48

    async def analyze_expressions(
        self,
        expression_paths: List[str],
    ) -> List[str]:
        if not expression_paths:
            return []

        collage_path = None
        docx_path = None
        try:
            collage_path = self._create_collage(expression_paths)
            docx_path = self._create_docx_with_collage(collage_path, len(expression_paths))

            somark_result = await somark_request_manager.parse_document(
                file_path=docx_path,
            )

            raw_sections = somark_result.get("sections", [])
            logger.info(f"SoMark 表情理解返回: code={somark_result.get('code')}, sections={len(raw_sections)}")
            if raw_sections:
                for i, sec in enumerate(raw_sections):
                    sec_type = sec.get("type", "?")
                    sec_text = sec.get("text", "")[:100]
                    has_image = "image" in sec or sec.get("type") == "image"
                    logger.info(f"  section[{i}]: type={sec_type}, has_image={has_image}, text_preview={sec_text!r}")
            else:
                logger.info(f"SoMark 表情理解原始返回全文(前500字): {str(somark_result)[:500]}")

            descriptions = self._parse_somark_result(somark_result, len(expression_paths))
            return descriptions

        except Exception as e:
            logger.error(f"表情分析失败: {e}", exc_info=True)
            return []
        finally:
            for path in [collage_path, docx_path]:
                if path and os.path.exists(path):
                    try:
                        os.remove(path)
                    except Exception:
                        pass

    def _create_collage(self, image_paths: List[str]) -> str:
        images = []
        for path in image_paths:
            try:
                img = Image.open(path).convert("RGBA")
                images.append(img)
            except Exception as e:
                logger.warning(f"加载表情图片失败 {path}: {e}")

        if not images:
            raise ValueError("没有可合成的表情图片")

        max_h = max(img.height for img in images)
        total_w = sum(img.width for img in images)
        spacing = self.COLLAGE_PADDING * (len(images) + 1)
        canvas_w = total_w + spacing
        canvas_h = max_h + self.COLLAGE_PADDING * 2 + self.LABEL_HEIGHT

        canvas = Image.new("RGBA", (canvas_w, canvas_h), (255, 255, 255, 255))
        draw = ImageDraw.Draw(canvas)

        try:
            font = ImageFont.truetype("simhei.ttf", self.FONT_SIZE)
        except Exception:
            try:
                font = ImageFont.truetype("msyh.ttf", self.FONT_SIZE)
            except Exception:
                font = ImageFont.load_default()

        x = self.COLLAGE_PADDING
        for i, img in enumerate(images):
            y = self.COLLAGE_PADDING
            canvas.paste(img, (x, y), img if img.mode == "RGBA" else None)

            label = f"expression_{i}"
            bbox = draw.textbbox((0, 0), label, font=font)
            label_w = bbox[2] - bbox[0]
            label_x = x + (img.width - label_w) // 2
            label_y = y + img.height + 10
            draw.text((label_x, label_y), label, fill=(0, 0, 0), font=font)

            x += img.width + self.COLLAGE_PADDING

        tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
        tmp_path = tmp.name
        canvas.save(tmp_path, format="PNG", optimize=True)
        tmp.close()
        logger.info(f"表情拼图已创建: {tmp_path} ({canvas_w}x{canvas_h})")
        return tmp_path

    def _create_docx_with_collage(self, collage_path: str, num_expressions: int) -> str:
        try:
            from docx import Document
            from docx.shared import Inches, Pt
        except ImportError:
            logger.warning("python-docx 不兼容，回退到直接使用图片")
            return collage_path

        doc = Document()
        section = doc.sections[0]
        section.page_width = Inches(28)
        section.page_height = Inches(16)
        section.top_margin = Inches(0.3)
        section.bottom_margin = Inches(0.3)
        section.left_margin = Inches(0.3)
        section.right_margin = Inches(0.3)

        p = doc.add_paragraph()
        p.alignment = 1
        run = p.add_run("Expression Diff Images")
        run.font.size = Pt(14)
        run.bold = True

        for i in range(num_expressions):
            p = doc.add_paragraph()
            run = p.add_run(f"expression_{i}:")
            run.bold = True
            run.font.size = Pt(10)

        p = doc.add_paragraph()
        p.alignment = 1
        run = p.add_run()
        run.add_picture(collage_path, width=Inches(24))

        tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        tmp_path = tmp.name
        doc.save(tmp_path)
        tmp.close()
        logger.info(f"表情DOCX已创建: {tmp_path}")
        return tmp_path

    def _parse_somark_result(self, result: Dict[str, Any], expected_count: int) -> List[str]:
        descriptions: List[str] = []

        raw_text = str(result)
        lines = raw_text.replace("\\n", "\n").split("\n")
        current_idx = -1
        buffer = ""

        for line in lines:
            for i in range(expected_count):
                if f"expression_{i}" in line:
                    if current_idx >= 0 and buffer:
                        while len(descriptions) <= current_idx:
                            descriptions.append("")
                        descriptions[current_idx] = buffer.strip()
                    current_idx = i
                    buffer = ""
                    break
            else:
                if current_idx >= 0:
                    buffer += " " + line.strip()

        if current_idx >= 0 and buffer:
            while len(descriptions) <= current_idx:
                descriptions.append("")
            descriptions[current_idx] = buffer.strip()
        logger.info(f"  文本行解析: {len(descriptions)} 个描述")
        for i, d in enumerate(descriptions):
            logger.info(f"    文本行[{i}]: {d[:150]!r}")

        sections = result.get("sections", [])
        section_texts = []
        for sec in sections:
            text = sec.get("text", "")
            if text:
                idx = -1
                for i in range(expected_count):
                    if f"expression_{i}" in text:
                        idx = i
                        text = text.replace(f"expression_{i}:", "").replace(f"expression_{i}", "").strip()
                        break
                if idx >= 0:
                    while len(section_texts) <= idx:
                        section_texts.append("")
                    if text:
                        section_texts[idx] = (section_texts[idx] + " " + text).strip()
                elif sec.get("image"):
                    img_text = text or self._describe_image_section(sec)
                    if img_text:
                        section_texts.append(f"[图像理解] {img_text}")
        logger.info(f"  sections解析: {len(section_texts)} 个描述")
        for i, d in enumerate(section_texts):
            logger.info(f"    sections[{i}]: {d[:150]!r}")

        result_descriptions = descriptions or section_texts

        while len(result_descriptions) < expected_count:
            result_descriptions.append(f"表情 {len(result_descriptions)}")

        result_descriptions = result_descriptions[:expected_count]
        result_descriptions = [d[:200] if len(d) > 200 else d for d in result_descriptions]

        logger.info(f"表情分析结果: {len(result_descriptions)} 个描述词")
        for i, d in enumerate(result_descriptions):
            logger.info(f"  表情 {i}: {d[:80]}...")
        return result_descriptions

    def _describe_image_section(self, sec: Dict[str, Any]) -> str:
        alt = sec.get("alt", "")
        text = sec.get("text", "")
        if alt and alt != text:
            return alt
        return text


expression_service = ExpressionService()
